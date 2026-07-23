import os
import PyPDF2
import docx
from typing import List
from backend.core.interfaces import BaseDocumentLoader
from backend.models.domain import Document
from backend.core.exceptions import DocumentLoadError
from backend.logging.logger import doc_logger

class LocalDocumentLoader(BaseDocumentLoader):
    """
    Loads documents from the local filesystem (txt, pdf, docx).
    """

    def load(self, source: str) -> List[Document]:
        doc_logger.info(f"Attempting to load document from: {source}")
        
        if not os.path.exists(source):
            raise DocumentLoadError(f"File not found: {source}")
            
        _, ext = os.path.splitext(source)
        ext = ext.lower()
        
        try:
            if ext == '.txt' or ext == '.md':
                content = self._load_txt(source)
            elif ext == '.pdf':
                content = self._load_pdf(source)
            elif ext == '.docx':
                content = self._load_docx(source)
            else:
                raise DocumentLoadError(f"Unsupported file extension: {ext}")
                
            doc = Document(
                content=content,
                source=source,
                metadata={"extension": ext, "filename": os.path.basename(source)}
            )
            doc_logger.info(f"Successfully loaded document: {source} ({len(content)} characters)")
            return [doc]
            
        except Exception as e:
            doc_logger.error(f"Failed to load document {source}: {str(e)}")
            raise DocumentLoadError(f"Error parsing document {source}: {str(e)}")

    def _load_txt(self, file_path: str) -> str:
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()

    def _load_pdf(self, file_path: str) -> str:
        text = ""
        with open(file_path, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\\n"
        return text

    def _load_docx(self, file_path: str) -> str:
        doc = docx.Document(file_path)
        return "\\n".join([para.text for para in doc.paragraphs])
